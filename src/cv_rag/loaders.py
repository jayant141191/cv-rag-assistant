"""Read CV files (PDF, DOCX, TXT, MD) into plain text."""

from __future__ import annotations

import io
import re
from pathlib import Path

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


def _read_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _read_docx(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    # CVs often put content in tables (two-column layouts).
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _read_text(data: bytes) -> str:
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def clean_text(text: str) -> str:
    text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def load_bytes(data: bytes, filename: str) -> str:
    """Extract clean text from raw file bytes; `filename` decides the parser."""
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type '{ext}'. Supported: {sorted(SUPPORTED_EXTENSIONS)}")
    if ext == ".pdf":
        text = _read_pdf(data)
    elif ext == ".docx":
        text = _read_docx(data)
    else:
        text = _read_text(data)
    text = clean_text(text)
    if not text:
        raise ValueError(
            f"No text could be extracted from '{filename}'. "
            "If it is a scanned PDF, run OCR on it first."
        )
    return text


def load_file(path: str | Path) -> str:
    path = Path(path)
    return load_bytes(path.read_bytes(), path.name)
